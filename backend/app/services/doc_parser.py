"""Turn a file from the document library (data/) into text for chunking.

Supports the formats the demo library uses: pdf, docx, xlsx, pptx, md/txt.
parse_structured() returns (line, section) pairs, where section is the nearest
heading / slide title / sheet name the line belongs to, so the chunker can
respect section boundaries and label every chunk with where it came from.
parse_file() is the same content as plain text.

Table/sheet structure is flattened to "col: value" style lines so a chunk
that is only a slice of a spreadsheet still carries its own headers.
"""

from pathlib import Path

Line = tuple[str, str]  # (text, section)


def parse_file(path: str | Path) -> str:
    return "\n".join(text for text, _ in parse_structured(path))


def is_table_format(path: str | Path) -> bool:
    return Path(path).suffix.lower() == ".xlsx"


def parse_structured(path: str | Path) -> list[Line]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in {".md", ".txt"}:
        return _parse_md(path)
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    if ext == ".xlsx":
        return _parse_xlsx(path)
    if ext == ".pptx":
        return _parse_pptx(path)
    raise ValueError(f"unsupported document format: {ext}")


def _parse_md(path: Path) -> list[Line]:
    out, section = [], ""
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            section = line.lstrip("#").strip()
        out.append((line, section))
    return out


def _parse_pdf(path: Path) -> list[Line]:
    from pypdf import PdfReader

    # No reliable heading signal in extracted PDF text: section stays empty.
    return [
        (line.strip(), "")
        for page in PdfReader(str(path)).pages
        for line in (page.extract_text() or "").splitlines()
        if line.strip()
    ]


def _parse_docx(path: Path) -> list[Line]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out: list[Line] = []
    section = ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                section = text
            out.append((text, section))
        elif tag == "tbl":
            table = Table(child, doc)
            header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
            for row in table.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                line = "; ".join(f"{h}: {v}" for h, v in zip(header, cells) if v)
                if line:
                    out.append((line, section))
    return out


def _parse_xlsx(path: Path) -> list[Line]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    out: list[Line] = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        section = f"Sheet: {ws.title}"
        out.append((f"[{section}]", section))
        header = [str(c) if c is not None else "" for c in rows[0]]
        if len(rows) == 1 or ws.title == "About":
            out += [(" | ".join(str(c) for c in r if c is not None), section) for r in rows]
            continue
        for r in rows[1:]:
            out.append(("; ".join(f"{h}: {c}" for h, c in zip(header, r) if c is not None and h), section))
    return out


def _parse_pptx(path: Path) -> list[Line]:
    from pptx import Presentation

    out: list[Line] = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        section = f"Slide {i}" + (f": {title}" if title else "")
        out.append((f"[Slide {i}]", section))
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append((shape.text_frame.text.strip(), section))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(("Notes: " + slide.notes_slide.notes_text_frame.text.strip(), section))
    return out
