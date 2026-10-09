"""Section-aware parsing and chunking (synthetic files, no external services)."""

from docx import Document as Docx
from openpyxl import Workbook

from app.services.doc_parser import is_table_format, parse_file, parse_structured
from app.services.ingestion import chunk_structured


def test_markdown_sections_follow_headings(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("# Title\n\nintro\n\n## Escalation\n\nSEV1 acks in 15 minutes\n", encoding="utf-8")
    assert parse_structured(f) == [("# Title", "Title"), ("intro", "Title"), ("## Escalation", "Escalation"),
                                   ("SEV1 acks in 15 minutes", "Escalation")]


def test_docx_in_document_order_with_heading_sections_and_table_rows(tmp_path):
    f = tmp_path / "a.docx"
    d = Docx()
    d.add_heading("Policy", 0)
    d.add_heading("Limits", 1)
    d.add_paragraph("Intro paragraph")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Band", "Multiplier"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "B", "0.8"
    d.add_heading("Review", 1)
    d.add_paragraph("Annual review")
    d.save(f)
    f_lines = parse_structured(f)
    assert ("Band: B; Multiplier: 0.8", "Limits") in f_lines
    assert [s for t_, s in f_lines if t_ == "Annual review"] == ["Review"]
    # table row sits between the paragraph before it and the next heading (document order)
    texts = [t_ for t_, _ in f_lines]
    assert texts.index("Intro paragraph") < texts.index("Band: B; Multiplier: 0.8") < texts.index("Review")


def test_xlsx_rows_carry_headers_and_sheet_section(tmp_path):
    f = tmp_path / "a.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "P&L"
    ws.append(["SGD million", "Q3"]); ws.append(["Total revenue", 18.4])
    wb.save(f)
    lines = parse_structured(f)
    assert ("SGD million: Total revenue; Q3: 18.4", "Sheet: P&L") in lines
    assert is_table_format(f) and "Total revenue" in parse_file(f)


def test_chunks_are_prefixed_and_cut_at_sheet_boundaries_for_tables():
    lines = [("[Sheet: A]", "Sheet: A"), ("row a1", "Sheet: A"), ("[Sheet: B]", "Sheet: B"), ("row b1", "Sheet: B")]
    chunks = chunk_structured(lines, title="Accounts", department="Finance", table=True)
    assert chunks == ["[Accounts | Finance | Sheet: A]\n[Sheet: A]\nrow a1",
                      "[Accounts | Finance | Sheet: B]\n[Sheet: B]\nrow b1"]


def test_prose_does_not_cut_tiny_sections_but_does_cut_big_ones():
    small = [("a" * 50, "S1"), ("b" * 50, "S2")]
    assert len(chunk_structured(small, title="T")) == 1  # under MIN_SECTION_CHARS: carried over
    big = [("a" * 400, "S1"), ("b" * 50, "S2")]
    out = chunk_structured(big, title="T")
    assert [c.split("\n")[0] for c in out] == ["[T | S1]", "[T | S2]"]


def test_oversized_line_is_window_split_and_size_limits_hold():
    out = chunk_structured([("x" * 2000, "")], title="T")
    assert len(out) >= 3 and all(len(c) <= 800 + len("[T]\n") for c in out)
    rows = [(f"row {i}: " + "v" * 100, "Sheet: S") for i in range(30)]
    tchunks = chunk_structured(rows, title="T", table=True)
    assert len(tchunks) > 1 and all(len(c) <= 1500 + 40 for c in tchunks)
